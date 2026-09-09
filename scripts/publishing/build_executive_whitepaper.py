#!/usr/bin/env python3
"""Build editable executive whitepaper and readable Markdown from shared content.

Run with the Python executable supplied by the artifact workspace runtime.
"""
import argparse,json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def link(paragraph,text,url):
    element=OxmlElement('w:hyperlink');element.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');prop=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'164C63');prop.append(color)
    run.append(prop);content=OxmlElement('w:t');content.text=text;run.append(content);element.append(run);paragraph._p.append(element)


def table(doc,values):
    columns=len(values[0])
    if columns not in (2,3) or any(len(row)!=columns for row in values):
        raise ValueError('Expected a rectangular two- or three-column table')
    result=doc.add_table(rows=0,cols=columns);result.alignment=WD_TABLE_ALIGNMENT.CENTER;result.autofit=False
    widths=[Cm(5.1),Cm(11.5)] if columns==2 else [Cm(4.2),Cm(5.8),Cm(6.6)]
    for col,width in zip(result.columns,widths):col.width=width
    for i,row in enumerate(values):
        cells=result.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=widths[j];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.08
            r=p.add_run(text);r.font.size=Pt(10.5);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
            props=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'16394A' if i==0 else ('F0F4F5' if i%2==0 else 'FFFFFF'));props.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');margins.append(el)
            props.append(margins);borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            props.append(borders)
        if i==0:
            repeat=OxmlElement('w:tblHeader');result.rows[i]._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');result.rows[i]._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--content',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();data=json.loads(args.content.read_text());args.output.mkdir(parents=True,exist_ok=True)
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Header','Footer']:
        style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
        rpr=style.element.get_or_add_rPr();fonts=OxmlElement('w:rFonts')
        for kind in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+kind),'Arial')
        rpr.append(fonts)
    normal=doc.styles['Normal'];normal.font.size=Pt(11.5);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(25);doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(21);doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 1'].paragraph_format.space_after=Pt(16)
    for root in [doc.styles.element,doc.element]:
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Footer'].font.size=Pt(8)
    header=sec.header.paragraphs[0];header.text='Engineering Governance    Geschäftsführung';header.style=doc.styles['Header'];header.runs[0].font.size=Pt(8)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run(f'Stand {data["date"]}    ');r.font.size=Pt(8)
    for text,field in [('Seite ','PAGE'),(' von ','NUMPAGES')]:
        footer.add_run(text).font.size=Pt(8);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),field);footer._p.append(f)
    md=[f"# {data['title']}",f"{data['subtitle']} · Ausgabe {data['version']} · {data['date']}",f"Quellstand: `{data['source_commit']}`"]
    for index,page in enumerate(data['pages']):
        if index:doc.add_page_break()
        doc.add_paragraph(page['title'],'Title' if index==0 else 'Heading 1')
        if index==0:
            doc.add_paragraph(data['subtitle'],'Subtitle')
            meta=doc.add_paragraph(f"Ausgabe {data['version']}\n{data['date']}\nProjekt devsecops-governance-framework")
            for run in meta.runs:run.font.size=Pt(10)
        else:md.append('## '+page['title'])
        for text in page['paragraphs']:doc.add_paragraph(text);md.append(text)
        if 'table' in page:
            table(doc,page['table']);md.append('\n'.join(['| '+' | '.join(page['table'][0])+' |','|'+'|'.join(['---']*len(page['table'][0]))+'|']+['| '+' | '.join(r)+' |' for r in page['table'][1:]]))
        if index==len(data['pages'])-1:
            doc.add_paragraph('Quellen zum festgehaltenen Stand','Heading 2');md.append('### Quellen zum festgehaltenen Stand')
            for source in data['sources']:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1
                p.add_run('['+source['id']+'] ').font.size=Pt(9)
                link(p,source['title'],source['url'])
                for run in p.runs:run.font.size=Pt(9)
                # Set hyperlink size explicitly as it is not in paragraph.runs.
                for run in p._p.findall('.//'+qn('w:r')):
                    props=run.find(qn('w:rPr'))
                    if props is None:props=OxmlElement('w:rPr');run.insert(0,props)
                    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');props.append(size)
                md.append(f"- [{source['id']}] [{source['title']}]({source['url']})")
    doc.core_properties.title=data['title'];doc.core_properties.subject=data['subtitle'];doc.core_properties.author='Projekt devsecops-governance-framework';doc.core_properties.keywords='Engineering Governance, Geschäftsführung, Testbetrieb'
    doc.save(args.output/'engineering-governance-whitepaper.docx')
    (args.content.parent/'whitepaper.md').write_text('\n\n'.join(md)+'\n')
    slide_md=['# Engineering Governance Präsentation für die Geschäftsführung',f"Quellstand: `{data['source_commit']}`. Vorgesehene Vortragsdauer etwa 15 Minuten."]
    for i,s in enumerate(data['slides'],1):
        slide_md += [f"## Folie {i} {s['title']}",s['lead']]
        slide_md += [f"- **{title}:** {body}" for title,body in s['items']]
        if 'table' in s:slide_md.append('\n'.join(['| '+' | '.join(s['table'][0])+' |','|---|---|']+['| '+' | '.join(r)+' |' for r in s['table'][1:]]))
        slide_md += ['### Sprechernotizen',s['notes'],'Quellen: '+', '.join('['+r+']('+next(x['url'] for x in data['sources'] if x['id']==r)+')' for r in s['refs'])]
    (args.content.parent/'presentation.md').write_text('\n\n'.join(slide_md)+'\n')
    print(args.output/'engineering-governance-whitepaper.docx')

if __name__=='__main__':main()
